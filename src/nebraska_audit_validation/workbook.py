"""Build the independent validation workbook from versioned audit records."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Iterable, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _cell_value(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_table(sheet, rows: Sequence[Mapping], headers: Sequence[str] | None = None) -> None:
    headers = list(headers or (rows[0].keys() if rows else []))
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append([_cell_value(row.get(header, "")) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, 1):
        values = [str(header)] + [str(row.get(header, "")) for row in rows[:250]]
        width = min(max(len(value) for value in values) + 2, 55)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build(repo_root: Path, output: Path) -> None:
    reports = repo_root / "reports"
    records = repo_root / "records"
    config = repo_root / "config"
    quality = json.loads((reports / "phase3_initial_quality.json").read_text())
    baseline = json.loads((reports / "baseline_metrics.json").read_text())
    resolutions = json.loads((config / "resolved_findings.json").read_text())["resolutions"]

    book = Workbook()
    readme = book.active
    readme.title = "Read Me"
    readme_rows = [
        {"Field": "Purpose", "Value": "Independent validation results; source observations are not overwritten."},
        {"Field": "Source repository", "Value": "trltrl12/Nebraska-Audit"},
        {"Field": "Source commit", "Value": baseline["scope"]["source_commit"]},
        {"Field": "Validation repository", "Value": "trltrl12/Nebraska-Audit-Validation"},
        {"Field": "Active failures", "Value": quality["severity_counts"].get("FAIL", 0)},
        {"Field": "Incomplete checks", "Value": quality["severity_counts"].get("INCOMPLETE", 0)},
        {"Field": "Active review candidates", "Value": quality["severity_counts"].get("REVIEW", 0)},
        {"Field": "Publication rule", "Value": "Only evidence-reviewed records that pass all applicable publication gates are verified."},
    ]
    _write_table(readme, readme_rows)

    _write_table(book.create_sheet("Investigations"), _csv_rows(records / "investigation_log.csv"))
    _write_table(book.create_sheet("Active Findings"), quality["findings"])
    _write_table(book.create_sheet("Priority Queue"), quality["priority_queue"])
    _write_table(book.create_sheet("Resolved Flags"), resolutions)
    _write_table(book.create_sheet("Correction Log"), _csv_rows(records / "correction_log.csv"), [
        "entity", "fiscal_year", "original_value", "corrected_value", "affected_metric",
        "reason", "source_evidence", "code_or_rule_change", "validation_result",
    ])
    inventory = [{"Metric": key, "Value": value} for key, value in baseline["inventory"].items()]
    _write_table(book.create_sheet("Baseline Inventory"), inventory)

    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=Path("."))
    parser.add_argument("--output", type=Path, default=Path("reports/nebraska_audit_validation.xlsx"))
    args = parser.parse_args(argv)
    build(args.repo.resolve(), args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
